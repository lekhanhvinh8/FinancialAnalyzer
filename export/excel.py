"""
Export P/E snapshot to Excel for a given index.

Usage:
    python export/excel.py vn30
    python export/excel.py vndiamond
    python export/excel.py vnfinlead
    python export/excel.py vn30 2026-07-01   # optional as-of date
"""

import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import os
from datetime import date as dt_date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db
from calculate import _quarter_end

_INDEX_LABELS = {
    "vn30":      "VN30",
    "vndiamond": "VN Diamond",
    "vnfinlead": "VN Fin Lead",
}

import functools
_MEMBER_LOADERS = {
    "vn30":      functools.partial(db.load_index_members, "vn30"),
    "vndiamond": functools.partial(db.load_index_members, "vndiamond"),
    "vnfinlead": functools.partial(db.load_index_members, "vnfinlead"),
}

# ── Styles ────────────────────────────────────────────────────────────────────

_FONT   = "Arial"
_BLUE   = Font(name=_FONT, color="0000FF", size=10)          # hardcoded inputs
_BLACK  = Font(name=_FONT, color="000000", size=10)          # formulas
_BOLD   = Font(name=_FONT, bold=True, size=10)
_TITLE  = Font(name=_FONT, bold=True, size=12)
_RED    = Font(name=_FONT, color="CC0000", bold=True, size=10)
_GREEN  = Font(name=_FONT, color="006400", bold=True, size=10)
_PURPLE = Font(name=_FONT, color="7030A0", bold=True, size=10)

_FILL_HEADER   = PatternFill("solid", fgColor="D9E1F2")
_FILL_14       = PatternFill("solid", fgColor="DDEEFF")
_FILL_13       = PatternFill("solid", fgColor="E2EFDA")
_FILL_INFO     = PatternFill("solid", fgColor="FFF2CC")
_FILL_TOTAL    = PatternFill("solid", fgColor="F2F2F2")

_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_RIGHT  = Alignment(horizontal="right",  vertical="center")
_LEFT   = Alignment(horizontal="left",   vertical="center")

_THIN  = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_FMT_VND   = '#,##0'
_FMT_PE    = '0.00"x"'
_FMT_PCT   = '0.00%'
_FMT_PRICE = '#,##0'


def _cell(ws, row, col, value=None, font=None, fill=None, align=None, fmt=None, border=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:   c.font      = font
    if fill:   c.fill      = fill
    if align:  c.alignment = align
    if fmt:    c.number_format = fmt
    if border: c.border    = border
    return c


def _get_symbol_data(symbol: str, as_of: pd.Timestamp):
    """Return dict with price, shares, 4 TTM quarters, or None if insufficient data."""
    profit_df = db.load_net_profit(symbol)
    price_df  = db.load_daily_prices(symbol)
    shares_df = db.load_shares(symbol)

    if profit_df.empty or price_df.empty or shares_df.empty:
        return None

    # Latest price ≤ as_of
    prices = price_df[price_df.index <= as_of]
    if prices.empty:
        return None
    price = float(prices["close"].iloc[-1])

    # Latest shares ≤ as_of
    shares_df["qend"] = shares_df.apply(
        lambda r: _quarter_end(int(r.year), int(r.quarter)), axis=1
    )
    s = shares_df[shares_df["qend"] <= as_of].sort_values("qend")
    if s.empty:
        return None
    shares = float(s.iloc[-1]["shares"])

    # 4 most recent quarters ≤ as_of
    profit_df["qend"] = profit_df.apply(
        lambda r: _quarter_end(int(r.year), int(r.quarter)), axis=1
    )
    q = profit_df[profit_df["qend"] <= as_of].sort_values("qend")
    if len(q) < 4:
        return None
    ttm = q.tail(4)

    return {"price": price, "shares": shares, "quarters": ttm}


def _q_val(quarters, col, idx):
    """Get float value or None from a quarters DataFrame."""
    v = quarters.iloc[idx][col]
    return float(v) if pd.notna(v) else None


def _range_label(quarters) -> str:
    o, l = quarters.iloc[0], quarters.iloc[-1]
    return f"Q{int(o.quarter)}-{int(o.year)} to Q{int(l.quarter)}-{int(l.year)}"


def _ttm10(row) -> float | str:
    """Sum kqkd_10 quarters; return "" if any quarter is missing."""
    vals = [row[f"q{k}_10"] for k in [1, 2, 3, 4]]
    if any(v is None for v in vals):
        return ""
    return sum(vals)


def build_excel(index: str, as_of: str | None = None, suffix: str = "") -> str:
    db.init_db()

    if index not in _INDEX_LABELS:
        raise ValueError(f"Unknown index '{index}'. Choose from: {list(_INDEX_LABELS)}")

    as_of_str = as_of or dt_date.today().isoformat()
    as_of_ts  = pd.Timestamp(as_of_str)
    label     = _INDEX_LABELS[index]

    members = _MEMBER_LOADERS[index](as_of=as_of_str)["symbol"].tolist()

    # ── Collect per-symbol data ───────────────────────────────────────────────
    rows = []
    for sym in sorted(members):
        d = _get_symbol_data(sym, as_of_ts)
        if d is None:
            print(f"  Skipping {sym} — insufficient data")
            continue
        q = d["quarters"]  # oldest → latest (iloc[0]=oldest, iloc[3]=latest)

        rows.append({
            "symbol": sym,
            "price":  d["price"],
            "shares": d["shares"],
            # kqkd_14 quarters (latest=idx3, oldest=idx0)
            "q1_14": _q_val(q, "kqkd_14", 3),
            "q2_14": _q_val(q, "kqkd_14", 2),
            "q3_14": _q_val(q, "kqkd_14", 1),
            "q4_14": _q_val(q, "kqkd_14", 0),
            # kqkd_13 quarters
            "q1_13": _q_val(q, "kqkd_13", 3),
            "q2_13": _q_val(q, "kqkd_13", 2),
            "q3_13": _q_val(q, "kqkd_13", 1),
            "q4_13": _q_val(q, "kqkd_13", 0),
            # kqkd_10 quarters
            "q1_10": _q_val(q, "kqkd_10", 3),
            "q2_10": _q_val(q, "kqkd_10", 2),
            "q3_10": _q_val(q, "kqkd_10", 1),
            "q4_10": _q_val(q, "kqkd_10", 0),
            "range": _range_label(q),
        })

    if not rows:
        raise ValueError(f"No data available for {label} as of {as_of_str}")

    # ── Build workbook ────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = f"{label} PE"

    n      = len(rows)
    R_DATA = 6                  # first data row
    R_LAST = R_DATA + n - 1    # last data row
    R_TOT  = R_LAST + 1        # TOTAL row
    R_PE14 = R_TOT  + 2        # Aggregate P/E (kqkd_14) row
    R_PE13 = R_PE14 + 1        # Aggregate P/E (kqkd_13) row
    R_PE10 = R_PE13 + 1        # Aggregate P/E (kqkd_10) row

    # Column indices
    cA,cB,cC,cD,cE,cF = 1,2,3,4,5,6          # Shared Inputs
    cG,cH,cI,cJ,cK,cL,cM = 7,8,9,10,11,12,13  # KQKD_14 section
    cN,cO,cP,cQ,cR,cS = 14,15,16,17,18,19      # KQKD_13 section
    cT,cU = 20,21                               # Info

    # ── Row 1: Title ──────────────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cU)
    _cell(ws, 1, 1, f"{label} — Aggregate P/E Calculation (KQKD_14 / KQKD_13 / KQKD_10)",
          font=_TITLE, align=_CENTER)

    # ── Row 2: Subtitle ───────────────────────────────────────────────────────
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cU)
    _cell(ws, 2, 1,
          f"As of: {as_of_str}    |    P/E col M = best available (14→13→10)"
          f"    |    Blue = KQKD_14 section    |    Green = KQKD_13 section"
          f"    |    Securities/insurance companies use KQKD_10 only",
          font=_BLACK, align=_CENTER)

    # ── Row 4: Section headers ────────────────────────────────────────────────
    sections = [
        (cA, cF,  "Shared Inputs",                                         _FILL_HEADER),
        (cG, cM,  "KQKD_14 Priority  (parent-company shareholders)",       _FILL_14),
        (cN, cS,  "KQKD_13 Priority  (consolidated incl. minority interest)", _FILL_13),
        (cT, cU,  "Info",                                                   _FILL_INFO),
    ]
    for c1, c2, text, fill in sections:
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        _cell(ws, 4, c1, text, font=_BOLD, fill=fill, align=_CENTER)

    # ── Row 5: Column headers ─────────────────────────────────────────────────
    col_headers = [
        (cA, "Symbol"),
        (cB, "Price\n(VND)"),
        (cC, "Shares"),
        (cD, "TTM KQKD_14\n(VND)"),
        (cE, "TTM KQKD_13\n(VND)"),
        (cF, "TTM KQKD_10\n(VND)"),
        (cG, "Q1 (latest)\n(VND)"),
        (cH, "Q2\n(VND)"),
        (cI, "Q3\n(VND)"),
        (cJ, "Q4 (oldest)\n(VND)"),
        (cK, "Net Profit\n=SUM(Q1:Q4)"),
        (cL, "Market Cap\n(VND)"),
        (cM, "P/E\n[best: 14→13→10]"),
        (cN, "Q1 (latest)\n(VND)"),
        (cO, "Q2\n(VND)"),
        (cP, "Q3\n(VND)"),
        (cQ, "Q4 (oldest)\n(VND)"),
        (cR, "Net Profit\n=SUM(Q1:Q4)"),
        (cS, "P/E\n[13 only]"),
        (cT, "TTM Range"),
        (cU, "Weight"),
    ]
    fills = {**{c: _FILL_HEADER for c in range(cA, cF+1)},
             **{c: _FILL_14     for c in range(cG, cM+1)},
             **{c: _FILL_13     for c in range(cN, cS+1)},
             **{c: _FILL_INFO   for c in range(cT, cU+1)}}
    for col, text in col_headers:
        _cell(ws, 5, col, text, font=_BOLD, fill=fills[col], align=_CENTER, border=_BORDER)

    ws.row_dimensions[5].height = 36

    # ── Data rows ─────────────────────────────────────────────────────────────
    def safe(v):
        return v if v is not None else ""

    f_ = get_column_letter(cF)

    for i, row in enumerate(rows):
        r = R_DATA + i
        g, h_, i_, j = get_column_letter(cG), get_column_letter(cH), get_column_letter(cI), get_column_letter(cJ)
        n_, o_, p_, q_ = get_column_letter(cN), get_column_letter(cO), get_column_letter(cP), get_column_letter(cQ)
        l_ = get_column_letter(cL)

        # Shared inputs (blue)
        _cell(ws, r, cA, row["symbol"],  font=_BLUE,  align=_LEFT)
        _cell(ws, r, cB, row["price"],   font=_BLUE,  align=_RIGHT, fmt=_FMT_PRICE)
        _cell(ws, r, cC, row["shares"],  font=_BLUE,  align=_RIGHT, fmt=_FMT_VND)
        _cell(ws, r, cD, f"=SUM({g}{r}:{j}{r})", font=_BLACK, align=_RIGHT, fmt=_FMT_VND)
        _cell(ws, r, cE, f"=SUM({n_}{r}:{q_}{r})", font=_BLACK, align=_RIGHT, fmt=_FMT_VND)
        _cell(ws, r, cF, _ttm10(row), font=_BLUE, align=_RIGHT, fmt=_FMT_VND)

        # KQKD_14 quarterly values (blue)
        for col, key in [(cG,"q1_14"),(cH,"q2_14"),(cI,"q3_14"),(cJ,"q4_14")]:
            _cell(ws, r, col, safe(row[key]), font=_BLUE, align=_RIGHT, fmt=_FMT_VND)

        # KQKD_14 formulas (black)
        k_ = get_column_letter(cK)
        _cell(ws, r, cK, f"=SUM({g}{r}:{j}{r})",  font=_BLACK, align=_RIGHT, fmt=_FMT_VND)
        _cell(ws, r, cL, f"=B{r}*C{r}",            font=_BLACK, align=_RIGHT, fmt=_FMT_VND)
        # Coalesced P/E: kqkd_14 → kqkd_13 → kqkd_10
        r_col_letter = get_column_letter(cR)
        _cell(ws, r, cM,
              f"=IF({k_}{r}>0,{l_}{r}/{k_}{r},"
              f"IF({r_col_letter}{r}>0,{l_}{r}/{r_col_letter}{r},"
              f"IF({f_}{r}>0,{l_}{r}/{f_}{r},\"N/A\")))",
              font=_BLACK, align=_RIGHT, fmt=_FMT_PE)

        # KQKD_13 quarterly values (blue)
        for col, key in [(cN,"q1_13"),(cO,"q2_13"),(cP,"q3_13"),(cQ,"q4_13")]:
            _cell(ws, r, col, safe(row[key]), font=_BLUE, align=_RIGHT, fmt=_FMT_VND)

        # KQKD_13 formulas (black)
        r_ = get_column_letter(cR)
        _cell(ws, r, cR, f"=SUM({n_}{r}:{q_}{r})",         font=_BLACK, align=_RIGHT, fmt=_FMT_VND)
        _cell(ws, r, cS, f"=IF({r_}{r}>0,{l_}{r}/{r_}{r},\"N/A\")", font=_BLACK, align=_RIGHT, fmt=_FMT_PE)

        # Info (black)
        _cell(ws, r, cT, row["range"], font=_BLACK, align=_CENTER)
        _cell(ws, r, cU, f"={l_}{r}/{l_}{R_TOT}", font=_BLACK, align=_RIGHT, fmt=_FMT_PCT)

    # ── TOTAL row ──────────────────────────────────────────────────────────────
    _cell(ws, R_TOT, cA, "TOTAL", font=_BOLD, fill=_FILL_TOTAL, align=_CENTER)
    k_col = get_column_letter(cK)
    l_col = get_column_letter(cL)
    r_col = get_column_letter(cR)
    f_col = get_column_letter(cF)
    _cell(ws, R_TOT, cK, f"=SUM({k_col}{R_DATA}:{k_col}{R_LAST})", font=_BOLD, fill=_FILL_TOTAL, fmt=_FMT_VND)
    _cell(ws, R_TOT, cL, f"=SUM({l_col}{R_DATA}:{l_col}{R_LAST})", font=_BOLD, fill=_FILL_TOTAL, fmt=_FMT_VND)
    _cell(ws, R_TOT, cR, f"=SUM({r_col}{R_DATA}:{r_col}{R_LAST})", font=_BOLD, fill=_FILL_TOTAL, fmt=_FMT_VND)
    _cell(ws, R_TOT, cF, f"=SUM({f_col}{R_DATA}:{f_col}{R_LAST})", font=_BOLD, fill=_FILL_TOTAL, fmt=_FMT_VND)

    # ── Aggregate P/E rows ─────────────────────────────────────────────────────
    ws.merge_cells(start_row=R_PE14, start_column=cA, end_row=R_PE14, end_column=cJ)
    _cell(ws, R_PE14, cA,
          f"Aggregate P/E (KQKD_14)  =  Total Market Cap / Total Net Profit [14]",
          font=_RED, align=_LEFT)
    _cell(ws, R_PE14, cK,
          f"=IF({k_col}{R_TOT}>0,{l_col}{R_TOT}/{k_col}{R_TOT},\"N/A\")",
          font=_RED, fmt=_FMT_PE)

    ws.merge_cells(start_row=R_PE13, start_column=cA, end_row=R_PE13, end_column=cJ)
    _cell(ws, R_PE13, cA,
          f"Aggregate P/E (KQKD_13)  =  Total Market Cap / Total Net Profit [13]",
          font=_GREEN, align=_LEFT)
    _cell(ws, R_PE13, cK,
          f"=IF({r_col}{R_TOT}>0,{l_col}{R_TOT}/{r_col}{R_TOT},\"N/A\")",
          font=_GREEN, fmt=_FMT_PE)

    ws.merge_cells(start_row=R_PE10, start_column=cA, end_row=R_PE10, end_column=cJ)
    _cell(ws, R_PE10, cA,
          f"Aggregate P/E (KQKD_10)  =  Total Market Cap / Total Net Profit [10]  "
          f"[securities/insurance companies only]",
          font=_PURPLE, align=_LEFT)
    _cell(ws, R_PE10, cK,
          f"=IF({f_col}{R_TOT}>0,{l_col}{R_TOT}/{f_col}{R_TOT},\"N/A\")",
          font=_PURPLE, fmt=_FMT_PE)

    # ── Column widths ──────────────────────────────────────────────────────────
    widths = {
        cA: 8, cB: 12, cC: 14, cD: 18, cE: 18, cF: 18,
        cG: 18, cH: 15, cI: 15, cJ: 15, cK: 18, cL: 18, cM: 14,
        cN: 18, cO: 15, cP: 15, cQ: 15, cR: 18, cS: 12,
        cT: 20, cU: 10,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "B6"

    # ── Save ──────────────────────────────────────────────────────────────────
    os.makedirs("exports", exist_ok=True)
    out = f"exports/{index}_pe_{as_of_str}{('_' + suffix) if suffix else ''}.xlsx"
    wb.save(out)
    print(f"Saved: {out}  ({n} symbols)")
    return out


if __name__ == "__main__":
    plain = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a.lstrip("-").split("=")[0]: (a.split("=", 1)[1] if "=" in a else "") for a in sys.argv[1:] if a.startswith("--")}
    idx    = plain[0] if plain else "vndiamond"
    as_of  = plain[1] if len(plain) > 1 else None
    suffix = flags.get("suffix", "")
    build_excel(idx, as_of, suffix=suffix)
